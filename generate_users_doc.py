import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("اطلاعات و سطوح دسترسی کاربران سامانه مدیریت پروژه")
title_run.font.name = "B Nazanin"
title_run.font.size = Pt(20)
title_run.bold = True
title_run.font.color.rgb = RGBColor(29, 78, 216) # Cobalt blue
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle_p = doc.add_paragraph()
sub_run = subtitle_p.add_run("Project Management System - User Credentials & Role-Based Access Control (RBAC)")
sub_run.font.size = Pt(11)
sub_run.font.color.rgb = RGBColor(100, 116, 139)
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# Intro
intro_p = doc.add_paragraph()
intro_run = intro_p.add_run("این مستند شامل مشخصات کامل حساب‌های کاربری پیش‌فرض تعریف‌شده در پایگاه داده سامانه جهت ورود، تست و بررسی سطوح مختلف دسترسی (مدیر ارشد، مدیر پروژه، عضو تیم و مشاهده‌گر) می‌باشد:")
intro_run.font.size = Pt(11)

# Table of users
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["نام کاربر", "پست الکترونیک (ایمیل)", "رمز عبور پیش‌فرض", "نقش سازمانی", "سطح دسترسی و اختیارات"]
hdr_cells = table.rows[0].cells
for idx, text in enumerate(headers):
    hdr_cells[idx].text = text
    shading = parse_xml(r'<w:shd {} w:fill="1D4ED8"/>'.format(nsdecls('w')))
    hdr_cells[idx]._tc.get_or_add_tcPr().append(shading)
    p = hdr_cells[idx].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        run.font.size = Pt(10)

users_data = [
    ("اباالفضل داریوشی (Abalfazl)", "abalfazl@example.com", "password123", "مدیر ارشد (Super Admin) / مالک", "دسترسی کامل مدیریتی: ایجاد و حذف پروژه‌ها، تخصیص نقش به تمام کاربران، تغییر تنظیمات، تایید نهایی تسک‌ها"),
    ("مریم راد (Maryam Rad)", "maryam@example.com", "password123", "مدیر پروژه (Project Manager)", "مدیریت تسک‌ها، تخصیص کار به افراد، تعریف اسپرینت‌ها، دعوت اعضای جدید و تعیین نقش در فضای کاری"),
    ("علی رضایی (Ali Rezaei)", "ali@example.com", "password123", "عضو تیم / توسعه‌دهنده (Member)", "ایجاد تسک جدید، تغییر وضعیت تسک‌های خود، افزودن زیرتسک، آپلود ضمیمه و ثبت کامنت در گفتگوها"),
    ("سارا تهرانی (Sara Tehrani)", "sara@example.com", "password123", "مشاهده‌گر (Viewer)", "فقط‌خواندنی: مشاهده پروژه‌ها، تقویم جلالی، تایم‌لاین و داشبوردها بدون امکان ویرایش، حذف یا تغییر نقش")
]

for row_data in users_data:
    row_cells = table.add_row().cells
    for i, val in enumerate(row_data):
        row_cells[i].text = val
        p = row_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [1, 2, 3] else WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            run.font.size = Pt(9.5)

# Formatting table borders
for row in table.rows:
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        tcPr.append(borders)

doc.add_paragraph()

# Security note
h2 = doc.add_paragraph()
h2_run = h2.add_run("نکات امنیتی و ثبت‌نام کاربران جدید")
h2_run.bold = True
h2_run.font.size = Pt(13)
h2_run.font.color.rgb = RGBColor(30, 41, 59)

notes = [
    "۱. هر کاربر جدید می‌تواند از طریق فرم ثبت‌نام (Sign Up) در صفحه اصلی برای خود حساب کاربری ایجاد کند.",
    "۲. رمزهای عبور با الگوریتم رمزنگاری استاندارد PBKDF2-HMAC-SHA256 هش شده و در دیتابیس با امنیت کامل ذخیره می‌شوند.",
    "۳. احراز هویت نشست‌ها با توکن‌های امن JWT انجام می‌شود.",
    "۴. تنها افراد دارای نقش مدیر ارشد یا مدیر پروژه می‌توانند نقش سایر کاربران را در پنجره «مدیریت اعضا و نقش‌ها» ارتقا دهند."
]

for n in notes:
    p = doc.add_paragraph()
    r = p.add_run(n)
    r.font.size = Pt(10.5)

out_path = os.path.join(os.path.dirname(__file__), "اطلاعات_کاربران_تستی_Project_Users.docx")
doc.save(out_path)
print(f"Word document saved to: {out_path}")
